"""Render DOCX to HTML/PDF/images and compare structure to source markdown."""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

import mammoth
from docx import Document
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parent
MD_PATH = ROOT / "HALFAPP_AGENT_ACTION_DIRECTIVES.md"
DOCX_PATH = ROOT / "HALFAPP_AGENT_ACTION_DIRECTIVES.docx"
OUT_DIR = ROOT / ".verify"
HTML_PATH = OUT_DIR / "preview.html"
PDF_PATH = OUT_DIR / "preview.pdf"

PAGE_CSS = """
@page { size: letter; margin: 1in; }
body {
  font-family: Arial, Helvetica, sans-serif;
  font-size: 12pt;
  line-height: 1.35;
  color: #111;
  max-width: 6.5in;
  margin: 0 auto;
  padding: 24px;
}
h1 { font-size: 16pt; margin: 18pt 0 12pt; }
h2 { font-size: 14pt; margin: 16pt 0 10pt; }
h3 { font-size: 12pt; margin: 14pt 0 8pt; font-weight: bold; }
p { margin: 0 0 8pt; }
ul, ol { margin: 0 0 8pt 1.2em; padding: 0; }
li { margin: 0 0 4pt; }
blockquote {
  margin: 8pt 0 8pt 24pt;
  font-style: italic;
  color: #333;
}
code, pre { font-family: Consolas, monospace; font-size: 10pt; }
table { border-collapse: collapse; width: 100%; margin: 8pt 0; }
td, th { border: 1px solid #ccc; padding: 6px 8px; vertical-align: top; }
"""


def md_headings(md: str) -> list[tuple[int, str]]:
    out: list[tuple[int, str]] = []
    for line in md.splitlines():
        if line.startswith("### "):
            out.append((3, line[4:].strip()))
        elif line.startswith("## "):
            out.append((2, line[3:].strip()))
        elif line.startswith("# "):
            out.append((1, line[2:].strip()))
    return out


def docx_headings(doc: Document) -> list[tuple[int, str]]:
    out: list[tuple[int, str]] = []
    for p in doc.paragraphs:
        if p.style and p.style.name and p.style.name.startswith("Heading"):
            level = int(p.style.name.split()[-1])
            text = p.text.strip()
            if text:
                out.append((level, text))
    return out


def md_bullet_count(md: str) -> int:
    return sum(1 for line in md.splitlines() if line.startswith("- "))


def docx_bullet_count(doc: Document) -> int:
    return sum(
        1
        for p in doc.paragraphs
        if p.style and p.style.name in ("List Bullet", "List Paragraph")
        and p.text.strip()
    )


def render_assets() -> list[Path]:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    with DOCX_PATH.open("rb") as f:
        result = mammoth.convert_to_html(f)
    html = f"""<!DOCTYPE html>
<html><head><meta charset="utf-8"><style>{PAGE_CSS}</style></head>
<body>{result.value}</body></html>"""
    HTML_PATH.write_text(html, encoding="utf-8")

    images: list[Path] = []
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={"width": 816, "height": 1056})
        page.goto(HTML_PATH.resolve().as_uri())
        page.wait_for_load_state("networkidle")
        PDF_PATH.write_bytes(
            page.pdf(
                format="Letter",
                margin={"top": "1in", "right": "1in", "bottom": "1in", "left": "1in"},
            )
        )
        full = OUT_DIR / "full.png"
        page.screenshot(path=str(full), full_page=True)
        images.append(full)
        browser.close()

    try:
        import fitz  # PyMuPDF

        doc = fitz.open(PDF_PATH)
        for i in range(doc.page_count):
            pix = doc.load_page(i).get_pixmap(matrix=fitz.Matrix(2, 2))
            path = OUT_DIR / f"pdf-page-{i + 1:02d}.png"
            pix.save(path)
            images.append(path)
        doc.close()
    except ImportError:
        pass
    return images


def structural_checks() -> dict:
    md = MD_PATH.read_text(encoding="utf-8")
    doc = Document(DOCX_PATH)
    md_h = md_headings(md)
    dx_h = docx_headings(doc)
    checks = {
        "headings_match": md_h == dx_h,
        "md_heading_count": len(md_h),
        "docx_heading_count": len(dx_h),
        "md_bullets": md_bullet_count(md),
        "docx_bullets": docx_bullet_count(doc),
        "has_tables_in_docx": False,
    }
    with zipfile.ZipFile(DOCX_PATH) as zf:
        xml = zf.read("word/document.xml").decode("utf-8")
        checks["has_tables_in_docx"] = "<w:tbl" in xml
    checks["bullet_delta"] = abs(checks["md_bullets"] - checks["docx_bullets"])
    return checks


def main() -> None:
    checks = structural_checks()
    images = render_assets()
    print("STRUCTURAL", checks)
    print("RENDERED", PDF_PATH, *images, sep="\n")


if __name__ == "__main__":
    main()
