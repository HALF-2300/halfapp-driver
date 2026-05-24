"""Generate HALFAPP_AGENT_ACTION_DIRECTIVES.docx from the markdown source."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
ROOT = Path(__file__).resolve().parent
MD_PATH = ROOT / "HALFAPP_AGENT_ACTION_DIRECTIVES.md"
OUT_PATH = ROOT / "HALFAPP_AGENT_ACTION_DIRECTIVES.docx"


def set_document_defaults(doc: Document) -> None:
    section = doc.sections[0]
    section.page_height = Inches(11)
    section.page_width = Inches(8.5)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")

    for level, size in [(1, 16), (2, 14), (3, 12)]:
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Arial"
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(size)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")


def add_inline_runs(paragraph, text: str) -> None:
    parts = re.split(r"(`[^`]+`)", text)
    for part in parts:
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        else:
            paragraph.add_run(part)


def add_bullet(doc: Document, text: str, level: int = 0) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
    add_inline_runs(p, text)


def add_numbered(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    add_inline_runs(p, text)


def add_blockquote(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.italic = True
    run.font.name = "Arial"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(51, 51, 51)


def parse_markdown(md: str, doc: Document) -> None:
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]

        if not line.strip():
            i += 1
            continue

        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
            i += 1
            continue

        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
            i += 1
            continue

        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
            i += 1
            continue

        if line.startswith("> "):
            doc.add_paragraph()
            add_blockquote(doc, line[2:].strip())
            i += 1
            continue

        if re.match(r"^\d+\.\s", line):
            add_numbered(doc, re.sub(r"^\d+\.\s+", "", line))
            i += 1
            continue

        if line.startswith("- "):
            add_bullet(doc, line[2:].strip())
            i += 1
            continue

        if line.strip() in ("Tasks:", "Done when:", "Goal:", "Status:", "Current implementation:"):
            p = doc.add_paragraph()
            run = p.add_run(line.strip())
            run.bold = True
            run.font.name = "Arial"
            i += 1
            continue

        para_lines = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].startswith((
            "#", "- ", "> ", "Tasks:", "Done when:", "Goal:", "Status:", "Current implementation:"
        )) and not re.match(r"^\d+\.\s", lines[i]):
            if lines[i].startswith("## ") or lines[i].startswith("### "):
                break
            para_lines.append(lines[i])
            i += 1
        p = doc.add_paragraph()
        add_inline_runs(p, " ".join(l.strip() for l in para_lines))


def main() -> None:
    md = MD_PATH.read_text(encoding="utf-8")
    doc = Document()
    set_document_defaults(doc)
    parse_markdown(md, doc)
    doc.save(OUT_PATH)
    print(f"Wrote {OUT_PATH}")


if __name__ == "__main__":
    main()
